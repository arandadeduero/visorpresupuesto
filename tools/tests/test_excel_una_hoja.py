"""Los Excel publicados llevan una sola hoja (la del ejercicio): sin copias, borradores ni hojas de trabajo."""
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[2]
EXCELES = sorted(p for p in ROOT.glob("[0-9][0-9][0-9][0-9]/*.xlsx") if not p.name.startswith("~$"))


def test_hay_excel_de_los_tres_ejercicios():
    assert {p.parent.name for p in EXCELES} >= {"2024", "2025", "2026"}


@pytest.mark.parametrize("excel", EXCELES, ids=lambda p: f"{p.parent.name}/{p.name}")
def test_cada_excel_tiene_una_sola_hoja(excel):
    wb = openpyxl.load_workbook(excel, read_only=True)
    assert len(wb.sheetnames) == 1, f"{excel.name}: {wb.sheetnames}"
    wb.close()
